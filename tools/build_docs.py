"""Script document for fact checking, as Word (.docx) and print-ready HTML (printed to PDF).

Everything comes from content/script.json (lines, on-screen text, visuals, claims),
content/timeline.json (timecodes) and evidence/manifest.json (source snapshots), so the
document always says exactly what the video says.

  python3 tools/build_docs.py      writes docs/script/*.docx and build/script.html
  node tools/print_pdf.cjs         prints build/script.html to docs/script/*.pdf
"""

import html
import json
import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
NAME = "From-Farmland-to-Token-Script"
INK = RGBColor(0x34, 0x37, 0x40)
MUTED = RGBColor(0x5E, 0x62, 0x70)
TEAL = RGBColor(0x00, 0x9C, 0x9C)


def tc(t):
    m, s = divmod(t, 60)
    return f"{int(m)}:{s:04.1f}"


def load():
    script = json.load(open(os.path.join(ROOT, "content", "script.json")))
    timeline = json.load(open(os.path.join(ROOT, "content", "timeline.json")))
    manifest = {m["id"]: m for m in json.load(open(os.path.join(ROOT, "evidence", "manifest.json")))["sources"]}
    times = {s["id"]: s for s in timeline["scenes"]}
    return script, timeline, manifest, times


def used_sources(script):
    ids = []
    for sc in script["scenes"]:
        for c in sc["claims"]:
            if c["source"] not in ids:
                ids.append(c["source"])
    return ids


# ------------------------------------------------------------------ docx
def shade(cell, hex_fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcpr.append(shd)


def repeat_header(row):
    """Repeat a table's header row at the top of every page."""
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trpr.append(el)


def para(doc_or_cell, text, size=10, bold=False, color=INK, italic=False, space_after=4, font=None):
    p = doc_or_cell.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color
    if font:
        run.font.name = font
    p.paragraph_format.space_after = Pt(space_after)
    return p


def cell_text(cell, lines, size=9, bold=False, color=INK, font=None):
    cell.text = ""
    first = True
    for line in lines:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        r = p.add_run(line)
        r.font.size = Pt(size)
        r.bold = bold
        r.font.color.rgb = color
        if font:
            r.font.name = font
        p.paragraph_format.space_after = Pt(2)


def build_docx(script, timeline, manifest, times):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Readex Pro"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Readex Pro")
    st.font.size = Pt(10)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(1.8)
        s.top_margin = s.bottom_margin = Cm(1.6)
        s.page_width, s.page_height = Cm(29.7), Cm(21.0)  # A4 landscape, wide tables

    para(doc, "SCRIPT AND FACT CHECK", 9, True, TEAL, space_after=2)
    para(doc, script["title"], 26, True, INK, space_after=2)
    para(doc, script["subtitle"], 12, False, MUTED, space_after=10)

    meta = [
        ("Asset", script["asset"]["name"] + ". " + script["asset"]["type"]),
        ("Runtime", f"{timeline['duration']:.1f} seconds ({tc(timeline['duration'])}), 16:9, 1920 x 1080, 60 fps"),
        ("Audience", script["audience"]),
        ("Voiceover", f"{script['voice']['voice']} ({script['voice']['engine']}), {script['voice']['dataset']}"),
        ("Note", script["asset"]["note"]),
    ]
    tbl = doc.add_table(rows=0, cols=2)
    for k, v in meta:
        row = tbl.add_row().cells
        cell_text(row[0], [k], 9, True, MUTED)
        cell_text(row[1], [v], 9)
        row[0].width, row[1].width = Cm(3.5), Cm(22)
    para(doc, "", space_after=4)
    para(doc, "Why this asset", 12, True, INK, space_after=2)
    para(doc, script["asset"]["why"], 10, space_after=10)

    para(doc, "1. Script", 16, True, INK, space_after=4)
    para(doc, "Timecodes are taken from the rendered voiceover. The on-screen column lists every word drawn in the frame during that scene.", 9, color=MUTED, space_after=6)
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [Cm(2.0), Cm(3.4), Cm(8.4), Cm(5.6), Cm(6.6)]
    repeat_header(table.rows[0])
    for c, h, w in zip(table.rows[0].cells, ["Time", "Scene", "Voiceover", "On-screen text", "Visual"], widths):
        cell_text(c, [h], 9, True, RGBColor(0xFF, 0xFF, 0xFF))
        shade(c, "343740")
        c.width = w
    for sc in script["scenes"]:
        t = times[sc["id"]]
        row = table.add_row().cells
        vo = [ln["vo"] for ln in sc["lines"]] or ["(no voiceover, end card holds)"]
        cell_text(row[0], [f"{tc(t['start'])} to {tc(t['end'])}"], 8, font="JetBrains Mono")
        cell_text(row[1], [sc["title"]], 9, True)
        cell_text(row[2], vo, 9)
        cell_text(row[3], sc["onscreen"], 8, color=MUTED)
        cell_text(row[4], [sc["visual"]], 8, color=MUTED)
        for c, w in zip(row, widths):
            c.width = w

    doc.add_page_break()
    para(doc, "2. Fact check", 16, True, INK, space_after=4)
    para(doc, "Every claim about Redbelly Network or Tokeniser, with the exact words of the supplied resource it comes from. Quotes are verbatim and are checked automatically against dated snapshots of each page (tools/check.py).", 9, color=MUTED, space_after=6)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    widths = [Cm(1.6), Cm(7.0), Cm(10.6), Cm(6.8)]
    repeat_header(table.rows[0])
    for c, h, w in zip(table.rows[0].cells, ["Scene", "Claim in the video", "Source quote (verbatim)", "Source"], widths):
        cell_text(c, [h], 9, True, RGBColor(0xFF, 0xFF, 0xFF))
        shade(c, "343740")
        c.width = w
    for sc in script["scenes"]:
        for cl in sc["claims"]:
            m = manifest[cl["source"]]
            row = table.add_row().cells
            cell_text(row[0], [sc["id"].upper()], 8, True, font="JetBrains Mono")
            cell_text(row[1], [cl["claim"]], 9)
            cell_text(row[2], [f"“{cl['quote']}”"], 9)
            cell_text(row[3], [m["url"], f"Snapshot {m['captured_utc'][:10]}"], 8, color=MUTED)
            for c, w in zip(row, widths):
                c.width = w

    para(doc, "", space_after=4)
    para(doc, "3. What the video does not say", 16, True, INK, space_after=4)
    for line in [
        "No prices, returns, yields, performance or investment advice. The fund is an illustrative example and no amounts are shown.",
        "No promise of liquidity. Tokeniser's own FAQ states that tokenisation does not guarantee liquidity, so the video does not mention liquidity at all.",
        "Atomic settlement is described only for investors who hold their own (non-custodial) wallets, as the Tokeniser knowledge hub specifies.",
        "No performance figures for Redbelly Network (throughput, uptime, validators) are used, so nothing in the video can age out of date.",
    ]:
        para(doc, "•  " + line, 10, space_after=3)

    para(doc, "", space_after=4)
    para(doc, "4. Sources", 16, True, INK, space_after=4)
    para(doc, "Text snapshots of every page are kept in the project files (evidence/) with their capture time and SHA-256 hash.", 9, color=MUTED, space_after=6)
    for sid in used_sources(script):
        m = manifest[sid]
        para(doc, f"{m['title']}", 10, True, space_after=0)
        para(doc, f"{m['url']}   captured {m['captured_utc'][:16].replace('T', ' ')} UTC   sha256 {m['sha256'][:16]}", 8, color=MUTED, space_after=5, font="JetBrains Mono")

    out = os.path.join(ROOT, "docs", "script", f"{NAME}.docx")
    os.makedirs(os.path.dirname(out), exist_ok=True)
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.title = f"{script['title']}: script and fact check"
    doc.save(out)
    return out


# ------------------------------------------------------------------ html for pdf
CSS = """
@font-face { font-family: 'Readex Pro'; src: url('../assets/fonts/ReadexPro-VF.woff2') format('woff2'); font-weight: 160 700; }
@font-face { font-family: 'JetBrains Mono'; src: url('../assets/fonts/JetBrainsMono-VF.woff2') format('woff2'); font-weight: 100 800; }
@page { size: A4 landscape; margin: 14mm 14mm 16mm; }
* { box-sizing: border-box; }
body { font-family: 'Readex Pro', sans-serif; color: #343740; font-size: 9.5pt; line-height: 1.45; margin: 0; }
.eyebrow { font-family: 'JetBrains Mono', monospace; color: #009C9C; letter-spacing: .14em; font-size: 8pt; font-weight: 500; text-transform: uppercase; }
h1 { font-size: 28pt; font-weight: 500; margin: 2mm 0 1mm; letter-spacing: -.01em; }
.sub { color: #5E6270; font-size: 12pt; margin: 0 0 6mm; }
h2 { font-size: 15pt; font-weight: 500; margin: 7mm 0 2mm; break-after: avoid; }
p.note { color: #5E6270; font-size: 8.5pt; margin: 0 0 3mm; }
.meta { display: grid; grid-template-columns: 30mm 1fr; gap: 1.5mm 6mm; margin: 0 0 5mm; padding: 4mm 5mm; background: #F6F5F2; border: 1px solid #E6E3DE; border-radius: 3mm; }
.meta b { color: #5E6270; font-weight: 500; }
table { width: 100%; border-collapse: collapse; font-size: 8.8pt; }
th { background: #343740; color: #fff; text-align: left; font-weight: 500; padding: 2mm 2.5mm; }
td { border-bottom: 1px solid #E6E3DE; padding: 2.2mm 2.5mm; vertical-align: top; }
tr { break-inside: avoid; }
td.tc, td.mono { font-family: 'JetBrains Mono', monospace; font-size: 7.8pt; color: #5E6270; white-space: nowrap; }
td.scene { font-weight: 500; }
td.vo p { margin: 0 0 1.5mm; }
td.on { color: #4B4F5D; font-size: 8pt; }
td.on span { display: inline-block; background: #ECEAE6; border-radius: 1.5mm; padding: .3mm 1.6mm; margin: 0 1mm 1mm 0; }
td.visual { color: #5E6270; font-size: 8pt; }
td.quote { font-size: 8.6pt; }
td.src a { color: #343740; word-break: break-all; font-size: 7.6pt; }
td.src small { color: #868382; display: block; font-size: 7.4pt; }
.step { display: inline-block; font-family: 'JetBrains Mono', monospace; font-size: 7pt; background: #A2FFC8; border-radius: 1mm; padding: 0 1.4mm; margin-bottom: 1mm; }
ul { margin: 0; padding-left: 5mm; }
li { margin-bottom: 1.5mm; }
.sources div { margin-bottom: 2.4mm; }
.sources .u { font-family: 'JetBrains Mono', monospace; font-size: 7.4pt; color: #5E6270; }
.foot { margin-top: 6mm; padding-top: 3mm; border-top: 1px solid #E6E3DE; color: #868382; font-size: 7.6pt; }
"""


def build_html(script, timeline, manifest, times):
    e = html.escape
    rows = []
    for sc in script["scenes"]:
        t = times[sc["id"]]
        vo = "".join(f"<p>{e(ln['vo'])}</p>" for ln in sc["lines"]) or "<p><i>No voiceover. End card holds.</i></p>"
        step = f'<span class="step">STEP {sc["step"]}</span><br>' if sc["step"] else ""
        on = "".join(f"<span>{e(x)}</span>" for x in sc["onscreen"])
        rows.append(
            f'<tr><td class="tc">{tc(t["start"])}<br>{tc(t["end"])}</td><td class="scene">{step}{e(sc["title"])}</td>'
            f'<td class="vo">{vo}</td><td class="on">{on}</td><td class="visual">{e(sc["visual"])}</td></tr>'
        )
    claims = []
    for sc in script["scenes"]:
        for cl in sc["claims"]:
            m = manifest[cl["source"]]
            claims.append(
                f'<tr><td class="mono">{sc["id"].upper()}</td><td>{e(cl["claim"])}</td>'
                f'<td class="quote">“{e(cl["quote"])}”</td>'
                f'<td class="src"><a href="{e(m["url"])}">{e(m["url"])}</a><small>Snapshot {m["captured_utc"][:10]}</small></td></tr>'
            )
    srcs = "".join(
        f'<div><b>{e(manifest[s]["title"])}</b><br><span class="u">{e(manifest[s]["url"])} · captured {manifest[s]["captured_utc"][:16].replace("T", " ")} UTC · sha256 {manifest[s]["sha256"][:16]}</span></div>'
        for s in used_sources(script)
    )
    a = script["asset"]
    nclaims = sum(len(sc["claims"]) for sc in script["scenes"])
    return f"""<!doctype html><html lang="en"><head><meta charset="utf-8"><title>{e(script['title'])}: script and fact check</title><style>{CSS}</style></head><body>
<div class="eyebrow">Script and fact check</div>
<h1>{e(script['title'])}</h1>
<p class="sub">{e(script['subtitle'])}</p>
<div class="meta">
<b>Asset</b><span>{e(a['name'])}. {e(a['type'])}</span>
<b>Why this asset</b><span>{e(a['why'])}</span>
<b>Runtime</b><span>{timeline['duration']:.1f} seconds ({tc(timeline['duration'])}). 16:9, 1920 x 1080, 60 fps, subtitles burned in and supplied as SRT and WebVTT.</span>
<b>Audience</b><span>{e(script['audience'])}</span>
<b>Voiceover</b><span>{e(script['voice']['voice'])} ({e(script['voice']['engine'])}). {e(script['voice']['dataset'])}.</span>
<b>Claims</b><span>{nclaims} claims about Redbelly Network or Tokeniser, each matched to a verbatim quote from the supplied resources (section 2).</span>
<b>Note</b><span>{e(a['note'])}</span>
</div>
<h2>1. Script</h2>
<p class="note">Timecodes come from the rendered voiceover. The on-screen column lists every word drawn in the frame during that scene.</p>
<table><thead><tr><th style="width:17mm">Time</th><th style="width:30mm">Scene</th><th style="width:78mm">Voiceover</th><th style="width:62mm">On-screen text</th><th>Visual</th></tr></thead><tbody>{''.join(rows)}</tbody></table>
<h2 style="break-before:page">2. Fact check</h2>
<p class="note">Every claim about Redbelly Network or Tokeniser, with the exact words of the supplied resource it comes from. Quotes are verbatim and are checked automatically against dated snapshots of each page.</p>
<table><thead><tr><th style="width:14mm">Scene</th><th style="width:70mm">Claim in the video</th><th>Source quote (verbatim)</th><th style="width:64mm">Source</th></tr></thead><tbody>{''.join(claims)}</tbody></table>
<h2>3. What the video does not say</h2>
<ul>
<li>No prices, returns, yields, performance or investment advice. The fund is an illustrative example and no amounts are shown.</li>
<li>No promise of liquidity. Tokeniser's own FAQ states that tokenisation does not guarantee liquidity, so the video does not mention liquidity at all.</li>
<li>Atomic settlement is described only for investors who hold their own (non-custodial) wallets, as the Tokeniser knowledge hub specifies.</li>
<li>No performance figures for Redbelly Network (throughput, uptime, validators) are used, so nothing in the video can age out of date.</li>
</ul>
<h2>4. Sources</h2>
<p class="note">Text snapshots of every page are kept in the project files (evidence/) with their capture time and SHA-256 hash.</p>
<div class="sources">{srcs}</div>
<div class="foot">Version {e(script['version'])}. Generated from content/script.json, content/timeline.json and evidence/manifest.json.</div>
</body></html>"""


def main():
    script, timeline, manifest, times = load()
    print("wrote", build_docx(script, timeline, manifest, times))
    os.makedirs(os.path.join(ROOT, "build"), exist_ok=True)
    out = os.path.join(ROOT, "build", "script.html")
    open(out, "w").write(build_html(script, timeline, manifest, times))
    print("wrote", out)


if __name__ == "__main__":
    main()
